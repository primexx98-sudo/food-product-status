"""에너지 카테고리 v2 검토용 — 전체 목록을 xlsx로 출력."""
import sys
import os
import glob

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))
sys.path.insert(0, os.path.dirname(__file__))

from category_mapper import categorize_general_food, is_general_excluded
from category_mapper_v2 import categorize_general_food_v2

from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(ROOT, 'output')


def load_all_items():
    items = {}
    for path in sorted(glob.glob(os.path.join(OUTPUT_DIR, '*.xlsx'))):
        wb = load_workbook(path, data_only=True)
        ws = wb['일반식품']
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        idx = {h: i for i, h in enumerate(header)}
        for r in rows[1:]:
            report_no = str(r[idx['품목제조번호']] or '')
            items[report_no] = (
                str(r[idx['품목명']] or ''),
                str(r[idx['원재료명']] or ''),
                str(r[idx['업소명']] or ''),
            )
    return items


def main():
    items = load_all_items()
    wb = Workbook()
    ws = wb.active
    ws.title = '에너지 v2 검토'
    ws.append(['품목명', '업소명', '기존 카테고리', 'v2 카테고리', '변경여부', '원재료명'])
    for report_no, (name, raw, company) in items.items():
        if is_general_excluded(name, report_no, company):
            continue
        old_cat = categorize_general_food(name, raw, report_no)
        new_cat = categorize_general_food_v2(name, raw, report_no)
        if old_cat != '에너지' and new_cat != '에너지':
            continue
        changed = '유지' if old_cat == new_cat else ('신규편입' if new_cat == '에너지' else '이탈')
        ws.append([name, company, old_cat, new_cat, changed, raw])
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)
    out_path = os.path.join(os.path.dirname(__file__), '에너지_v2_검토.xlsx')
    wb.save(out_path)
    print('저장:', out_path)


if __name__ == '__main__':
    main()
