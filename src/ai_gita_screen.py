"""'기타'로 분류된 일반식품 품목 중 건강기능식품/웰빙 트렌드와 명백히 무관한 것을
Gemini API로 스크리닝해 제외 후보 목록을 만든다. 코드/데이터를 자동으로 고치지 않고
`output/AI_제외후보.md`에 후보만 정리해둔다 — 실제 제외 규칙 반영은 사람이 검토 후 진행.

사용법: GEMINI_API_KEY 환경변수 설정 후 `python ai_gita_screen.py`
"""
import glob
import json
import os
import time

import requests
from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATE_PATH = os.path.join(ROOT_DIR, 'output', 'ai_screen_state.json')
CANDIDATES_PATH = os.path.join(ROOT_DIR, 'output', 'AI_제외후보.md')

GEMINI_MODEL = 'gemini-3.6-flash'
GEMINI_URL = f'https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent'
BATCH_SIZE = 25

# 이미 "기타 유지가 정책"으로 확정된 항목(설계서.md 7절) — AI가 이걸 제외 후보로
# 다시 올리지 않도록 프롬프트에 명시
_ALREADY_DECIDED_KEEP = (
    '생강', '알룰로스', '스테비아', '헛개나무', '흑마늘', '삼백초', '민들레',
    '젤리', '비트', '비트루트', '녹용',
)

PROMPT_TEMPLATE = """너는 건강기능식품/웰빙 트렌드 데이터를 정리하는 보조자다.
아래는 식품안전나라에 신고된 일반식품 품목 중 아직 어떤 기능성 카테고리에도 매칭되지 않아
"기타"로 남아있는 것들이다. 각 품목이 건강기능식품/웰빙 트렌드 데이터로서 의미가 있는지,
아니면 그냥 일반 식자재·조리식품·과자·B2B 원료라 트렌드 집계에서 제외하는 게 맞는지 판단해줘.

## 제외 후보로 볼 것
- 기능성 브랜딩/원료 없이 그냥 맛(과일맛/캐릭터/전통간식)만 있는 기호식품
- 조리용 반가공 식자재(볶음용/구이용/찌개용 등), B2B 원료·파우더·믹스류
- 냉동 완제품 간식, 베이커리/제과 부재료, 캐릭터 상품
- 전통 원물 단순가공품(썰기/말리기/굽기 정도, 기능성 브랜딩 없음)

## 유지할 것(제외 후보 아님)
- 원료명에 구체적 기능성 성분이 있는데 아직 카테고리 매칭 로직이 못 잡은 것
- 다음 키워드가 포함된 품목은 이미 "기타 유지"가 정책으로 확정돼 있으니 절대 제외 후보로
  올리지 말 것: {already_decided}
- 애매하면 반드시 유지(keep) 쪽으로 판단 — 오탐 방지가 최우선

## 출력 형식
반드시 JSON 배열만 출력. 각 원소는 {{"report_no": "...", "verdict": "keep" 또는
"exclude_candidate", "reason": "10단어 이내 한글 근거"}}. 입력에 없는 report_no를
만들어내지 말고, 입력된 모든 항목에 대해 정확히 하나씩만 응답할 것.

## 품목 목록 (품목제조번호 | 품목명 | 원재료명)
{items}
"""


def _load_all_gita_items():
    items = {}
    for path in sorted(glob.glob(os.path.join(ROOT_DIR, 'output', '*_품목신고보고현황.xlsx'))):
        wb = load_workbook(path, data_only=True)
        if '일반식품' not in wb.sheetnames:
            continue
        ws = wb['일반식품']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            report_no, prms_dt, lcns_no, name, shape, company, chng_dt, raw, cat = row[:9]
            if cat != '기타':
                continue
            report_no = str(report_no or '')
            if not report_no:
                continue
            items[report_no] = {
                'name': name or '',
                'raw': raw or '',
                'company': company or '',
            }
    return items


def _load_state():
    if os.path.exists(STATE_PATH):
        with open(STATE_PATH, encoding='utf-8') as f:
            return json.load(f)
    return {}


def _save_state(state):
    with open(STATE_PATH, 'w', encoding='utf-8') as f:
        json.dump(state, f, ensure_ascii=False, indent=2, sort_keys=True)


def _call_gemini(api_key, batch):
    lines = [f"{rno} | {info['name']} | {(info['raw'] or '')[:200]}" for rno, info in batch]
    prompt = PROMPT_TEMPLATE.format(
        already_decided=', '.join(_ALREADY_DECIDED_KEEP),
        items='\n'.join(lines),
    )
    body = {
        'contents': [{'parts': [{'text': prompt}]}],
        'generationConfig': {'responseMimeType': 'application/json'},
    }
    resp = requests.post(
        GEMINI_URL, params={'key': api_key}, json=body, timeout=60,
    )
    resp.raise_for_status()
    data = resp.json()
    text = data['candidates'][0]['content']['parts'][0]['text']
    return json.loads(text)


def _write_candidates_md(state):
    candidates = [
        (rno, info) for rno, info in state.items()
        if info.get('verdict') == 'exclude_candidate'
    ]
    candidates.sort(key=lambda x: x[1].get('name', ''))
    lines = [
        '# AI 제외 후보 목록',
        '',
        f'`ai_gita_screen.py`가 Gemini로 스크리닝한 결과. 자동 반영되지 않음 — 사람이 검토 후',
        '`src/category_mapper.py`의 `_GENERAL_EXCLUDE_*`에 직접 반영할 것.',
        '',
        f'총 {len(candidates)}건',
        '',
        '| 품목명 | 업소명 | 품목제조번호 | 근거 |',
        '|---|---|---|---|',
    ]
    for rno, info in candidates:
        lines.append(
            f"| {info.get('name','')} | {info.get('company','')} | {rno} | {info.get('reason','')} |"
        )
    with open(CANDIDATES_PATH, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')


def main():
    api_key = os.environ.get('GEMINI_API_KEY')
    if not api_key:
        print('[건너뜀] GEMINI_API_KEY 미설정')
        return

    all_items = _load_all_gita_items()
    state = _load_state()
    new_report_nos = [rno for rno in all_items if rno not in state]

    if not new_report_nos:
        print('신규 스크리닝 대상 없음')
        _write_candidates_md(state)
        return

    print(f'신규 스크리닝 대상: {len(new_report_nos)}건')
    for i in range(0, len(new_report_nos), BATCH_SIZE):
        batch_nos = new_report_nos[i:i + BATCH_SIZE]
        batch = [(rno, all_items[rno]) for rno in batch_nos]
        try:
            results = _call_gemini(api_key, batch)
        except Exception as e:
            print(f'  배치 {i}~{i+len(batch)} 실패: {e}')
            continue
        result_by_no = {str(r.get('report_no', '')): r for r in results}
        for rno, info in batch:
            r = result_by_no.get(rno)
            if not r:
                continue
            state[rno] = {
                'name': info['name'],
                'company': info['company'],
                'verdict': r.get('verdict', 'keep'),
                'reason': r.get('reason', ''),
            }
        time.sleep(1)

    _save_state(state)
    _write_candidates_md(state)
    n_candidates = sum(1 for v in state.values() if v.get('verdict') == 'exclude_candidate')
    print(f'완료. 누적 제외 후보 {n_candidates}건 -> {CANDIDATES_PATH}')


if __name__ == '__main__':
    main()
