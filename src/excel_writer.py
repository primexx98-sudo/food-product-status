import glob
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from category_mapper import HEALTH_CAT_LIST, GENERAL_CAT_LIST

HEALTH_COLUMNS = [
    '품목제조번호', '보고일자', '제품형태', '최종생성일시',
    '품목명', '업소명', '주된기능성', '소비기한', '섭취방법', '원재료', '카테고리',
]
GENERAL_COLUMNS = [
    '품목제조번호', '보고일자', '인허가번호', '품목명',
    '제품형태', '업소명', '최근수정일자', '원재료명', '카테고리',
]

_BLUE_FILL   = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
_LIGHT_FILL  = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
_WHITE_FONT  = Font(color='FFFFFF', bold=True)
_BOLD_FONT   = Font(bold=True)
_CENTER      = Alignment(horizontal='center', vertical='center')
_BLUE_FONT   = Font(color='4472C4')


def _autofit(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 55)


def _write_data_sheet(ws, data: list[dict], columns: list[str]):
    for ci, name in enumerate(columns, 1):
        c = ws.cell(row=1, column=ci, value=name)
        c.fill = _BLUE_FILL
        c.font = _WHITE_FONT
        c.alignment = _CENTER

    for ri, item in enumerate(data, 2):
        for ci, col in enumerate(columns, 1):
            ws.cell(row=ri, column=ci, value=item.get(col, ''))

    ws.freeze_panes = 'A2'
    _autofit(ws)


def _gather_monthly_counts(output_dir: str, year: int,
                           current_month: int,
                           current_health: list[dict],
                           current_general: list[dict]) -> tuple[dict, dict]:
    monthly_h = {m: {c: 0 for c in HEALTH_CAT_LIST}   for m in range(1, 13)}
    monthly_g = {m: {c: 0 for c in GENERAL_CAT_LIST}  for m in range(1, 13)}

    pattern = os.path.join(output_dir, f"{year}-*_품목신고보고현황.xlsx")
    for fpath in sorted(glob.glob(pattern)):
        fname = os.path.basename(fpath)
        try:
            file_month = int(fname[5:7])
        except (ValueError, IndexError):
            continue
        if file_month == current_month:
            continue

        try:
            wb = load_workbook(fpath, read_only=True, data_only=True)
            for sheet_name, monthly_dict, cat_list in [
                ('건강기능식품', monthly_h, HEALTH_CAT_LIST),
                ('일반식품',     monthly_g, GENERAL_CAT_LIST),
            ]:
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                rows = ws.iter_rows(values_only=True)
                headers = list(next(rows, []))
                if '카테고리' not in headers:
                    continue
                cat_idx = headers.index('카테고리')
                for row in rows:
                    cat = row[cat_idx] if row[cat_idx] else '기타'
                    if cat in monthly_dict[file_month]:
                        monthly_dict[file_month][cat] += 1
            wb.close()
        except Exception as e:
            print(f"Warning: {fpath} 읽기 실패 — {e}")

    for item in current_health:
        cat = item.get('카테고리', '기타')
        if cat in monthly_h[current_month]:
            monthly_h[current_month][cat] += 1

    for item in current_general:
        cat = item.get('카테고리', '기타')
        if cat in monthly_g[current_month]:
            monthly_g[current_month][cat] += 1

    return monthly_h, monthly_g


def _write_summary_table(ws, title: str, monthly_data: dict,
                         cat_list: list[str], row_start: int):
    # 제목
    title_cell = ws.cell(row=row_start, column=2, value=title)
    title_cell.font = Font(bold=True, size=13, color='1F3864')

    header_row = row_start + 1
    ws.cell(row=header_row, column=1, value='').fill = _LIGHT_FILL
    for ci, cat in enumerate(cat_list, 2):
        c = ws.cell(row=header_row, column=ci, value=cat)
        c.fill = _LIGHT_FILL
        c.font = _BOLD_FONT
        c.alignment = _CENTER

    total_col = len(cat_list) + 2
    tc = ws.cell(row=header_row, column=total_col, value='합계')
    tc.fill = _LIGHT_FILL
    tc.font = _BOLD_FONT
    tc.alignment = _CENTER

    for month in range(1, 13):
        dr = header_row + month
        ws.cell(row=dr, column=1, value=f"{month:02d}월").font = _BOLD_FONT
        row_total = 0
        for ci, cat in enumerate(cat_list, 2):
            count = monthly_data[month].get(cat, 0)
            c = ws.cell(row=dr, column=ci, value=count)
            c.alignment = _CENTER
            if count == 0:
                c.font = _BLUE_FONT
            row_total += count
        total_c = ws.cell(row=dr, column=total_col, value=row_total)
        total_c.alignment = _CENTER
        total_c.font = _BOLD_FONT


def _write_dashboard(ws, output_dir: str, year: int,
                     current_month: int,
                     health_data: list[dict],
                     general_data: list[dict]):
    monthly_h, monthly_g = _gather_monthly_counts(
        output_dir, year, current_month, health_data, general_data
    )

    # A1 = 연도
    year_cell = ws.cell(row=1, column=1, value=year)
    year_cell.font = Font(bold=True, size=14)

    _write_summary_table(ws, '건강기능식품 품목제조신고 현황 - 기능성 별',
                         monthly_h, HEALTH_CAT_LIST, row_start=1)

    # 건강기능식품 테이블: rows 1~14 (title + header + 12 months)
    general_start = 1 + 14 + 2  # = 17
    _write_summary_table(ws, '일반식품 품목제조보고 현황 - 기능성 별',
                         monthly_g, GENERAL_CAT_LIST, row_start=general_start)

    # 컬럼 너비 조정
    ws.column_dimensions['A'].width = 8
    for ci in range(2, max(len(HEALTH_CAT_LIST), len(GENERAL_CAT_LIST)) + 4):
        ws.column_dimensions[get_column_letter(ci)].width = 10


def write_excel(filepath: str, output_dir: str, year: int, month: int,
                health_data: list[dict], general_data: list[dict]):
    wb = Workbook()
    wb.remove(wb.active)

    ws_h = wb.create_sheet('건강기능식품')
    ws_g = wb.create_sheet('일반식품')
    ws_d = wb.create_sheet('DASHBOARD')

    _write_data_sheet(ws_h, health_data, HEALTH_COLUMNS)
    _write_data_sheet(ws_g, general_data, GENERAL_COLUMNS)
    _write_dashboard(ws_d, output_dir, year, month, health_data, general_data)

    wb.save(filepath)
    print(f"저장 완료: {filepath}")
