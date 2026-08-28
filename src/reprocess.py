"""기존 수집된 엑셀 파일에 대해 category_mapper 로직(제외 규칙·카테고리)을 재적용.
API 재호출 없이 이미 저장된 원본 데이터만 다시 분류/필터링해서 덮어쓴다.
사용법: python reprocess.py 2026 6 [2026 7 ...]
"""
import sys
import os

from category_mapper import categorize_health_food, categorize_general_food, is_general_excluded
from excel_writer import write_excel, HEALTH_COLUMNS, GENERAL_COLUMNS

from openpyxl import load_workbook


def _read_sheet(filepath, sheet_name, columns):
    wb = load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {col: row[idx[col]] if col in idx and row[idx[col]] is not None else '' for col in columns}
        rows.append(item)
    return rows


def reprocess(root_dir, year, month):
    output_dir = os.path.join(root_dir, 'output')
    filename = f"{year}-{month:02d}_품목신고보고현황.xlsx"
    filepath = os.path.join(output_dir, filename)
    if not os.path.exists(filepath):
        print(f"[건너뜀] {filename} 없음")
        return

    health_data = _read_sheet(filepath, '건강기능식품', HEALTH_COLUMNS)
    for item in health_data:
        item['카테고리'] = categorize_health_food(
            item.get('주된기능성', ''), item.get('품목명', ''), item.get('원재료', ''),
        )

    general_data = _read_sheet(filepath, '일반식품', GENERAL_COLUMNS)
    before = len(general_data)
    general_data = [
        r for r in general_data
        if not is_general_excluded(r.get('품목명', ''), r.get('품목제조번호', ''), r.get('업소명', ''))
    ]
    after = len(general_data)
    general_data.sort(key=lambda x: str(x.get('보고일자', '') or ''))
    for item in general_data:
        item['카테고리'] = categorize_general_food(
            item.get('품목명', ''), item.get('원재료명', ''), item.get('품목제조번호', ''),
        )

    write_excel(filepath, output_dir, year, month, health_data, general_data)
    print(f"  일반식품 {before}건 -> {after}건 (제외 {before - after}건)")


if __name__ == '__main__':
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    args = sys.argv[1:]
    pairs = [(int(args[i]), int(args[i + 1])) for i in range(0, len(args), 2)]
    for year, month in pairs:
        print(f"[{year}-{month:02d}] 재분류 시작")
        reprocess(root_dir, year, month)
